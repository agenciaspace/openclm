import math
import re
from datetime import date
from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt
from fastapi import HTTPException
from pydantic import EmailStr, TypeAdapter, ValidationError

from .models import Contract, DocumentVersion
from .security import digest


def validate_answers(questions: list[dict], answers: dict):
    errors = {}
    if set(answers) - {q["key"] for q in questions}:
        errors["_form"] = "Há respostas para perguntas inexistentes."
    clean = {}
    for q in questions:
        key, kind = q["key"], q["type"]
        value = answers.get(key)
        if value is None or value == "":
            if q["required"]:
                errors[key] = "Campo obrigatório."
            else:
                clean[key] = None
            continue
        valid = True
        if kind == "boolean":
            valid = isinstance(value, bool)
        elif kind == "number":
            valid = type(value) in (int, float) and abs(value) <= 1e15 and math.isfinite(value)
        else:
            valid = isinstance(value, str) and bool(value.strip()) and len(value) <= 10000
            if valid:
                value = value.strip()
                if kind == "select":
                    valid = value in q["options"]
                elif kind == "date":
                    try:
                        valid = date.fromisoformat(value).isoformat() == value
                    except ValueError:
                        valid = False
                elif kind == "email":
                    try:
                        value = str(TypeAdapter(EmailStr).validate_python(value))
                    except ValidationError:
                        valid = False
        if not valid:
            errors[key] = f"Resposta inválida para o tipo {kind}."
        else:
            clean[key] = value
    if errors:
        raise HTTPException(
            422, {"message": "Revise as respostas do formulário.", "fields": errors}
        )
    return clean


def render(body: str, answers: dict):
    def substitute(match):
        value = answers.get(match.group(1))
        if isinstance(value, bool):
            return "Sim" if value else "Não"
        return str(value) if value is not None else "—"

    return re.sub(r"\{\{\s*([a-z][a-z0-9_]*)\s*\}\}", substitute, body)


def snapshot(db, contract: Contract):
    db.add(
        DocumentVersion(
            contract_id=contract.id,
            version=contract.version,
            title=contract.title,
            counterparty=contract.counterparty,
            content=contract.content,
            answers=contract.answers,
            sha256=digest(contract.content),
        )
    )


def docx_bytes(title: str, content: str, signer_count: int = 0, anchor_prefix: str = "openclm"):
    doc = Document()
    doc.sections[0].top_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, 0)
    for line in content.splitlines():
        doc.add_paragraph(line)
    if signer_count:
        doc.add_heading("Assinaturas", level=1)
        for index in range(1, signer_count + 1):
            doc.add_paragraph(f"Signatário {index}\n\n/{anchor_prefix}_sign_{index}/\n")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
